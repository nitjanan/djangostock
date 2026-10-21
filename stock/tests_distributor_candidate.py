from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from stock.models import (
    BaseBranchCompany, BaseVatType, BaseVisible, Distributor, DistributorCandidate, DistributorForm, DistributorRule,
    DistributorRuleGroup, UserProfile,
)
from stock.views import _annotate_system_status


OSM_SHOP = {
    "name": "ร้านเจริญการช่าง",
    "shop_type": "hardware",
    "address": "1 ถนนทดสอบ",
    "lat": 13.756331,
    "lng": 100.501765,
    "phone": "081-111-2222",
    "osm_id": 555,
    "osm_type": "node",
}

LEAF_CODES = ["1", "2", "3", "4", "5", "6", "7", "8", "9.1", "9.2", "10", "11"]


class DistributorCandidateTestCase(TestCase):
    """ร้านจากแผนที่ -> ผู้คัดเลือกกรอกใบ FM-PU-005 -> ส่ง -> กลุ่ม ApproveDistributor อนุมัติ -> เข้า Distributor"""

    def setUp(self):
        self.approver_group, _ = Group.objects.get_or_create(name='ApproveDistributor')
        self.user = User.objects.create_user(username='requester', password='password')
        self.other = User.objects.create_user(username='other', password='password')
        self.approver = User.objects.create_user(username='approver', password='password')
        self.outsider = User.objects.create_user(username='outsider', password='password')
        self.approver.groups.add(self.approver_group)
        branch = BaseBranchCompany.objects.create(id="1", code="HO", name="Head Office")
        for u in (self.user, self.other, self.approver, self.outsider):
            UserProfile.objects.create(user=u).branch_company.add(branch)
        # เมนู "ผู้จัดจำหน่าย": ผู้คัดเลือกเปิดแท็บ Distributor ใน UserProfile / ผู้อนุมัติเห็นครบจากกลุ่ม ApproveDistributor
        self.distributor_tab = BaseVisible.objects.get(name='Distributor')
        for u in (self.user, self.other):
            u.userprofile.visible.add(self.distributor_tab)
        self.rules = {r.code: r for r in DistributorRule.objects.all()}
        self.no_vat = BaseVatType.objects.create(id="0", name="ไม่มี vat")
        self.vat = BaseVatType.objects.create(id="1", name="รวม vat")

    def login(self, user):
        # RequireCompanyCodeMiddleware logout ทุกคนที่ session ไม่มี company_code
        self.client.force_login(user)
        session = self.client.session
        session["company_code"] = "HO"
        session.save()

    def propose(self, user=None, **overrides):
        data = {
            "place_source": "osm",
            "place_id": "node/555",
            "name": OSM_SHOP["name"],
            "address": OSM_SHOP["address"],
            "tel": OSM_SHOP["phone"],
            "shop_type": "ร้านวัสดุ",
            "lat": "13.756331",
            "lng": "100.501765",
        }
        data.update(overrides)
        self.login(user or self.user)
        return self.client.post(reverse('proposeDistributorCandidate'), data)

    def fill_form(self, candidate, action="save", answers=None, skip=(), **header):
        """ตอบผ่านทุกข้อ ยกเว้นที่ระบุใน answers / ข้ามข้อใน skip"""
        # หน้าฟอร์มส่งค่าส่วนหัวที่กรอกไว้แล้วกลับมาทุกช่อง
        data = {"action": action}
        for field in ("name", "address", "branch", "contact", "tel", "fax", "other_contact", "tax_id", "email"):
            data[field] = getattr(candidate, field) or ""
        data.update(header)
        answers = answers or {}
        for code in LEAF_CODES:
            if code not in skip:
                data[f"ans_{self.rules[code].pk}"] = answers.get(code, "pass")
        self.login(self.user)
        return self.client.post(reverse('distributorCandidateForm', args=[candidate.pk]), data)

    def submitted_candidate(self, answers=None, **header):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.fill_form(c, action="submit", answers=answers, **header)
        c.refresh_from_db()
        return c

    def approve(self, candidate, user=None, **data):
        if data.get("mode") == "new":
            data.setdefault("vat_type", self.no_vat.pk)
        self.login(user or self.approver)
        return self.client.post(reverse('approveDistributorCandidate', args=[candidate.pk]), data)

    # ----- หลักเกณฑ์จาก migration -----

    def test_rules_seeded_from_fm_pu_005(self):
        self.assertEqual(sorted(self.rules), sorted(LEAF_CODES + ["9"]))
        self.assertEqual(self.rules["9.1"].parent, self.rules["9"])
        self.assertEqual(DistributorRuleGroup.objects.count(), 4)
        self.assertEqual(self.rules["8"].mandatory_group, self.rules["9"].mandatory_group)
        self.assertIsNone(self.rules["3"].mandatory_group)

    # ----- ตรวจเกณฑ์บังคับ -----

    def evaluate_with(self, answers):
        c = DistributorCandidate.objects.create(place_source="osm", place_id=f"node/{DistributorCandidate.objects.count()}", name="x")
        for code in LEAF_CODES:
            DistributorForm.objects.create(candidate=c, question=self.rules[code], ans=answers.get(code, "pass"))
        return c.evaluate()

    def test_all_pass(self):
        e = self.evaluate_with({})
        self.assertTrue(e["mandatory_passed"])
        self.assertEqual(e["missing"], [])

    def test_mandatory_single_rule_fail(self):
        self.assertFalse(self.evaluate_with({"1": "fail"})["mandatory_passed"])
        # ข้อไม่บังคับไม่ผ่าน ไม่กระทบ
        self.assertTrue(self.evaluate_with({"3": "fail", "11": "fail"})["mandatory_passed"])

    def test_rule_8_or_9(self):
        self.assertTrue(self.evaluate_with({"8": "fail"})["mandatory_passed"])
        self.assertTrue(self.evaluate_with({"9.1": "fail", "9.2": "fail"})["mandatory_passed"])
        self.assertFalse(self.evaluate_with({"8": "fail", "9.1": "fail", "9.2": "fail"})["mandatory_passed"])

    def test_rule_9_children_mode(self):
        # ค่าเริ่มต้น all: 9.1 ไม่ผ่าน -> ข้อ 9 ไม่ผ่าน -> 8 ก็ไม่ผ่าน -> กลุ่มไม่ผ่าน
        self.assertFalse(self.evaluate_with({"8": "fail", "9.1": "fail"})["mandatory_passed"])
        DistributorRule.objects.filter(code="9").update(children_mode=DistributorRule.CHILDREN_ANY)
        self.assertTrue(self.evaluate_with({"8": "fail", "9.1": "fail"})["mandatory_passed"])

    def test_inactive_rule_not_required(self):
        DistributorRule.objects.filter(code="11").update(is_active=False)
        c = DistributorCandidate.objects.create(place_source="osm", place_id="node/1", name="x")
        for code in LEAF_CODES[:-1]:
            DistributorForm.objects.create(candidate=c, question=self.rules[code], ans="pass")
        e = c.evaluate()
        self.assertEqual(e["missing"], [])
        self.assertNotIn("11", [row["rule"].code for row in e["rows"]])

    # ----- สถานะในตารางผู้จัดจำหน่ายอื่นๆ -----

    def status_of(self, shop):
        vendors = [dict(shop)]
        _annotate_system_status(vendors, "osm")
        return vendors[0]

    def test_status_new_draft_pending_in_system(self):
        self.assertEqual(self.status_of(OSM_SHOP)["system_status"], "new")
        self.propose()
        self.assertEqual(self.status_of(OSM_SHOP)["system_status"], "draft")
        c = DistributorCandidate.objects.get()
        self.fill_form(c, action="submit")
        self.assertEqual(self.status_of(OSM_SHOP)["system_status"], "pending")
        self.approve(c, mode="new", distributor_id="V001")
        ev = self.status_of(OSM_SHOP)
        self.assertEqual((ev["system_status"], ev["distributor_id"]), ("in_system", "V001"))

    def test_osm_way_and_node_with_same_number_are_different_places(self):
        Distributor.objects.create(id="V001", name="x", place_source="osm", place_id="way/555")
        self.assertEqual(self.status_of(OSM_SHOP)["system_status"], "new")

    # ----- ยื่นฟอร์ม -----

    def test_propose_creates_draft_and_opens_form(self):
        response = self.propose()
        c = DistributorCandidate.objects.get()
        self.assertRedirects(response, reverse('distributorCandidateForm', args=[c.pk]), fetch_redirect_response=False)
        self.assertEqual(c.status, DistributorCandidate.STATUS_DRAFT)
        self.assertEqual(c.requested_by, self.user)
        self.assertEqual(c.latitude, Decimal("13.756331"))
        # กดซ้ำ -> กลับไปฟอร์มเดิม ไม่สร้างใบใหม่
        self.propose()
        self.assertEqual(DistributorCandidate.objects.count(), 1)

    def test_other_user_cannot_edit_draft(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        response = self.propose(user=self.other)
        self.assertRedirects(response, reverse('distributorCandidateDetail', args=[c.pk]), fetch_redirect_response=False)
        self.login(self.other)
        self.client.post(reverse('distributorCandidateForm', args=[c.pk]), {"action": "save", "name": "แก้ชื่อ"})
        c.refresh_from_db()
        self.assertEqual(c.name, OSM_SHOP["name"])

    def test_save_draft_stores_header_and_answers(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.fill_form(c, answers={"3": "fail"}, contact="คุณสมชาย", tax_id="0105555000000", fax="02-111-2222", other_contact="LINE: @shop\nFacebook: ร้านทดสอบ")
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_DRAFT)
        self.assertEqual((c.contact, c.tax_id, c.fax, c.other_contact), ("คุณสมชาย", "0105555000000", "02-111-2222", "LINE: @shop\nFacebook: ร้านทดสอบ"))
        self.assertEqual(c.answers.count(), len(LEAF_CODES))
        self.assertEqual(c.answers.get(question=self.rules["3"]).ans, "fail")

    def test_parent_rule_cannot_be_answered(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        self.client.post(reverse('distributorCandidateForm', args=[c.pk]), {
            "action": "save", "name": c.name, f"ans_{self.rules['9'].pk}": "pass",
        })
        self.assertFalse(c.answers.filter(question=self.rules["9"]).exists())

    def test_submit_requires_all_answers(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.fill_form(c, action="submit", skip=("5",))
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_DRAFT)

        self.fill_form(c, action="submit")
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertIsNotNone(c.submitted_at)

    def check(self, candidate, user=None, answers=None, skip=()):
        data = {}
        for code in LEAF_CODES:
            if code not in skip:
                data[f"ans_{self.rules[code].pk}"] = (answers or {}).get(code, "pass")
        self.login(user or self.user)
        return self.client.post(reverse('distributorCandidateCheck', args=[candidate.pk]), data)

    def test_check_button_evaluates_without_saving(self):
        self.propose()
        c = DistributorCandidate.objects.get()

        res = self.check(c).json()
        self.assertTrue(res["passed"])
        self.assertEqual(res["missing"], [])
        self.assertIn("สรุป: ผ่านเกณฑ์บังคับ", res["html"])
        self.assertFalse(c.answers.exists())  # ไม่บันทึก

        res = self.check(c, answers={"1": "fail"}).json()
        self.assertFalse(res["passed"])
        self.assertIn("ไม่ผ่านเกณฑ์บังคับ", res["html"])

        res = self.check(c, skip=("5", "9.2")).json()
        self.assertEqual(sorted(res["missing"]), sorted([self.rules["5"].pk, self.rules["9.2"].pk]))
        self.assertIn("ยังตอบไม่ครบ: ข้อ 5, 9.2", res["html"])

    def test_summary_explains_why_group_fails(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        # ข้อ 8 ไม่ผ่าน / 9.1 ไม่ผ่าน / 9.2 ยังไม่ตอบ -> กลุ่ม "ข้อ 8 หรือ ข้อ 9" ไม่ผ่าน
        res = self.check(c, answers={"8": "fail", "9.1": "fail"}, skip=("9.2", "4")).json()
        html = " ".join(res["html"].split())
        self.assertIn("(ผ่านข้อใดข้อหนึ่งก็พอ)", html)
        self.assertIn("ข้อ 8 <span class=\"text-danger\">ตอบว่าไม่ผ่าน</span>", html)
        self.assertIn("เพราะข้อย่อยต้องผ่านทั้งหมด", html)
        self.assertIn("ข้อ 9.1 ตอบว่าไม่ผ่าน", html)
        self.assertIn("ข้อ 9.2 ยังไม่ได้ตอบ", html)
        self.assertIn("ข้อ 4 <span class=\"text-secondary\">ยังไม่ได้ตอบ</span>", html)

        # ผ่านด้วยข้อ 8 อย่างเดียว -> บอกว่าผ่านจากข้อไหน
        res = self.check(c, answers={"9.1": "fail"}).json()
        self.assertIn("ผ่านจาก ข้อ 8", " ".join(res["html"].split()))

        evaluation = c.evaluate(answers={})
        group89 = next(g for g in evaluation["groups"] if len(g["members"]) == 2)
        nine = next(m for m in group89["members"] if m["rule"].code == "9")
        self.assertEqual((nine["state"], nine["need_all"]), ("parent", True))
        self.assertEqual([k["state"] for k in nine["children"]], ["missing", "missing"])

    def test_check_button_only_for_editor(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.assertEqual(self.check(c, user=self.other).status_code, 403)
        self.login(self.user)
        self.assertEqual(self.client.get(reverse('distributorCandidateCheck', args=[c.pk])).status_code, 405)
        self.assertContains(self.client.get(reverse('distributorCandidateForm', args=[c.pk])), "ตรวจเกณฑ์")

    def test_cancel_draft_deletes_it(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        self.client.post(reverse('distributorCandidateForm', args=[c.pk]), {"action": "cancel"})
        self.assertFalse(DistributorCandidate.objects.exists())

    # ----- อนุมัติ -----

    def test_approve_new_copies_form_into_distributor(self):
        c = self.submitted_candidate(contact="คุณสมชาย", fax="02-111-2222", other_contact="LINE: @shop", tax_id="0105555000000")
        self.approve(c, mode="new", distributor_id="V900")

        d = Distributor.objects.get(id="V900")
        self.assertEqual((d.name, d.address, d.tel), (OSM_SHOP["name"], OSM_SHOP["address"], OSM_SHOP["phone"]))
        self.assertEqual((d.contact, d.fax, d.tex), ("คุณสมชาย", "02-111-2222", "0105555000000"))
        self.assertEqual(d.other_contact, "LINE: @shop")
        self.assertEqual((d.place_source, d.place_id), ("osm", "node/555"))
        self.assertEqual((d.latitude, d.longitude), (Decimal("13.756331"), Decimal("100.501765")))
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_APPROVED)
        self.assertEqual((c.distributor, c.approved_by), (d, self.approver))
        self.assertIsNotNone(c.approved_at)

    def test_type_and_genre_dropdowns_saved_and_copied_on_approve(self):
        from stock.models import BaseDistributorGenre, BaseDistributorType
        transport = BaseDistributorType.objects.create(id="06", name="ขนส่ง")
        BaseDistributorType.objects.create(id="00", name="ไม่ทราบ")
        hauling = BaseDistributorGenre.objects.create(id="19", name="ค่าจ้างบรรทุก-บุคคลภายนอก")

        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        page = self.client.get(reverse('distributorCandidateForm', args=[c.pk]))
        self.assertContains(page, '<option value="06" >ขนส่ง</option>')
        self.assertContains(page, '<option value="19" >ค่าจ้างบรรทุก-บุคคลภายนอก</option>')

        self.fill_form(c, type="06", genre="19")
        c.refresh_from_db()
        self.assertEqual((c.type, c.genre), (transport, hauling))
        page = self.client.get(reverse('distributorCandidateForm', args=[c.pk]))
        self.assertContains(page, '<option value="06" selected>ขนส่ง</option>')

        # ค่าที่ไม่มีในตาราง = ไม่เลือก
        self.fill_form(c, type="99", genre="")
        c.refresh_from_db()
        self.assertEqual((c.type, c.genre), (None, None))

        self.fill_form(c, action="submit", type="06", genre="19")
        c.refresh_from_db()
        self.approve(c, mode="new", distributor_id="V900")
        d = Distributor.objects.get(id="V900")
        self.assertEqual((d.type, d.genre), (transport, hauling))
        self.login(self.approver)
        self.assertContains(self.client.get(reverse('distributorApproveDetail', args=[c.pk])), "ชนิดผู้จัดจำหน่าย: ขนส่ง")

    def test_approve_new_requires_vat_type(self):
        # ตาราง Distributor จริงตั้ง vat_type_id เป็น NOT NULL
        c = self.submitted_candidate()
        self.approve(c, mode="new", distributor_id="V900", vat_type="")
        self.assertFalse(Distributor.objects.filter(id="V900").exists())

        self.approve(c, mode="new", distributor_id="V900", vat_type=self.vat.pk)
        self.assertEqual(Distributor.objects.get(id="V900").vat_type, self.vat)

    def test_cannot_approve_when_mandatory_fails(self):
        c = self.submitted_candidate(answers={"2": "fail"})
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.approve(c, mode="new", distributor_id="V900")
        self.assertFalse(Distributor.objects.filter(id="V900").exists())
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)

    def test_cannot_approve_draft(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.fill_form(c)
        self.approve(c, mode="new", distributor_id="V900")
        self.assertFalse(Distributor.objects.filter(id="V900").exists())

    def test_approve_new_rejects_existing_express_id(self):
        Distributor.objects.create(id="V900", name="ร้านอื่น")
        c = self.submitted_candidate()
        self.approve(c, mode="new", distributor_id="V900")
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertIsNone(Distributor.objects.get(id="V900").place_id)

    def test_link_mode_removed_existing_id_is_not_linked(self):
        # หน้าอนุมัติไม่มีโหมด "ผูกกับผู้จัดจำหน่ายที่มีอยู่แล้ว" แล้ว: ส่ง mode=link มาก็ทำเหมือนสร้างใหม่ (รหัสซ้ำ -> ไม่ผ่าน)
        Distributor.objects.create(id="D-01", name="เจริญการช่าง")
        c = self.submitted_candidate()
        self.approve(c, mode="link", distributor_id="D-01")
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertIsNone(Distributor.objects.get(id="D-01").place_id)
        self.login(self.approver)
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertNotContains(response, "ผูกกับผู้จัดจำหน่ายที่มีอยู่แล้ว")
        self.assertNotContains(response, "ผูกรายเดิม")
        self.assertContains(response, "อนุมัติ + สร้างใหม่")

    def test_user_outside_group_cannot_approve_or_reject(self):
        c = self.submitted_candidate()
        self.approve(c, user=self.user, mode="new", distributor_id="V900")
        self.login(self.user)
        self.client.post(reverse('rejectDistributorCandidate', args=[c.pk]))
        self.assertFalse(Distributor.objects.filter(id="V900").exists())
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)

    def test_rejected_can_be_proposed_again_keeping_answers(self):
        c = self.submitted_candidate(answers={"3": "fail"})
        self.login(self.approver)
        self.client.post(reverse('rejectDistributorCandidate', args=[c.pk]), {"reject_reason": "ราคาสูง"})
        c.refresh_from_db()
        self.assertEqual((c.status, c.reject_reason, c.approved_by), (DistributorCandidate.STATUS_REJECTED, "ราคาสูง", self.approver))

        self.propose(user=self.other)
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_DRAFT)
        self.assertEqual(c.requested_by, self.other)
        self.assertIsNone(c.approved_by)
        # เหตุผลครั้งก่อนเก็บไว้ แสดงบนฟอร์มตอนแก้
        self.assertEqual(c.reject_reason, "ราคาสูง")
        self.assertEqual(c.answers.get(question=self.rules["3"]).ans, "fail")

    def reject(self, candidate, reason="ไม่มีหนังสือรับรอง"):
        self.login(self.approver)
        self.client.post(reverse('rejectDistributorCandidate', args=[candidate.pk]), {"reject_reason": reason})
        candidate.refresh_from_db()

    def test_reopen_button_on_rejected_tab(self):
        c = self.submitted_candidate(answers={"3": "fail"})
        self.reject(c)

        self.login(self.other)
        page = self.client.get(reverse('distributorCandidateList') + "?cand_status=rejected")
        self.assertContains(page, reverse('distributorCandidateReopen', args=[c.pk]))
        self.assertContains(page, "เหตุผล: ไม่มีหนังสือรับรอง")

        response = self.client.post(reverse('distributorCandidateReopen', args=[c.pk]))
        self.assertRedirects(response, reverse('distributorCandidateForm', args=[c.pk]), fetch_redirect_response=False)
        c.refresh_from_db()
        self.assertEqual((c.status, c.requested_by, c.approved_by), (DistributorCandidate.STATUS_DRAFT, self.other, None))
        self.assertEqual(c.answers.get(question=self.rules["3"]).ans, "fail")

        form = self.client.get(reverse('distributorCandidateForm', args=[c.pk]))
        self.assertContains(form, "ใบนี้เคยไม่อนุมัติ")
        self.assertContains(form, "ไม่มีหนังสือรับรอง")

        # ส่งใหม่แล้ว หน้าดูใบบอกว่าเป็นเหตุผล "ครั้งก่อน"
        self.login(self.other)
        self.client.post(reverse('distributorCandidateForm', args=[c.pk]), dict(
            {"action": "submit", "name": c.name}, **{f"ans_{self.rules[code].pk}": "pass" for code in LEAF_CODES}))
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertContains(self.client.get(reverse('distributorCandidateDetail', args=[c.pk])), "ครั้งก่อนไม่อนุมัติเพราะ: ไม่มีหนังสือรับรอง")

    def test_reopen_only_rejected_and_post(self):
        c = self.submitted_candidate()
        self.login(self.user)
        self.client.post(reverse('distributorCandidateReopen', args=[c.pk]))
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertEqual(self.client.get(reverse('distributorCandidateReopen', args=[c.pk])).status_code, 405)

        # ร้านเข้าระบบไปแล้ว (เช่นมีคนผูกไว้ทางอื่น) -> ไม่เปิดซ้ำ
        self.reject(c)
        Distributor.objects.create(id="V500", name="x", place_source="osm", place_id="node/555")
        self.login(self.user)
        self.client.post(reverse('distributorCandidateReopen', args=[c.pk]))
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_REJECTED)

    def test_propose_redirect_ignores_external_next(self):
        Distributor.objects.create(id="V001", name="x", place_source="osm", place_id="node/555")
        response = self.propose(next="https://evil.example.com/")
        self.assertEqual(response.url, reverse('viewExternalVendorReport'))

    # ----- หน้าจอ -----

    @patch("stock.views.fetch_vendors_from_openstreetmap")
    def test_report_filters_by_system_status(self, fetch):
        other = dict(OSM_SHOP, name="ร้านในระบบ", osm_id=777)
        fetch.side_effect = lambda *a, **k: [dict(OSM_SHOP), dict(other)]
        Distributor.objects.create(id="V777", name="ร้านในระบบ", place_source="osm", place_id="node/777")

        self.login(self.user)
        url = reverse('viewExternalVendorReport') + "?lat=13.75&lng=100.5&source=osm"
        response = self.client.get(url + "&system_status=new")
        self.assertEqual([ev["name"] for ev in response.context["external_vendors"]], [OSM_SHOP["name"]])
        self.assertContains(response, "เสนอเพิ่ม")

        response = self.client.get(url + "&system_status=in_system")
        self.assertEqual([ev["distributor_id"] for ev in response.context["external_vendors"]], ["V777"])
        self.assertContains(response, "มีในระบบ: V777")

    @patch("stock.views.fetch_vendors_from_tomtom")
    @patch("stock.views.fetch_vendors_from_openstreetmap")
    def test_report_lists_saved_candidates_without_map_search(self, fetch_osm, fetch_tomtom):
        self.propose()
        draft = DistributorCandidate.objects.get()
        approved = DistributorCandidate.objects.create(place_source="osm", place_id="node/9", name="ร้านอนุมัติแล้ว", status="approved")

        self.login(self.user)
        response = self.client.get(reverse('distributorCandidateList'))
        fetch_osm.assert_not_called()
        fetch_tomtom.assert_not_called()
        self.assertEqual([c.pk for c in response.context["candidates"]], [draft.pk])
        self.assertContains(response, reverse('distributorCandidateForm', args=[draft.pk]))

        response = self.client.get(reverse('distributorCandidateList') + "?cand_status=approved")
        self.assertEqual([c.pk for c in response.context["candidates"]], [approved.pk])

    def test_report_lists_newest_vendors_first(self):
        import datetime
        Distributor.objects.create(id="A01", name="เก่า", created=datetime.date(2024, 2, 16))
        Distributor.objects.create(id="A02", name="เก่า 2", created=datetime.date(2024, 2, 16))
        Distributor.objects.create(id="Z99", name="ใหม่", created=datetime.date(2026, 10, 1))
        self.login(self.user)
        response = self.client.get(reverse('viewVendorReport'))
        self.assertEqual([v.id for v in response.context["vender"]], ["Z99", "A02", "A01"])

    def test_report_shows_type_and_genre(self):
        from stock.models import BaseDistributorGenre, BaseDistributorType
        Distributor.objects.create(
            id="T01", name="ร้านมีชนิด",
            type=BaseDistributorType.objects.create(id="06", name="ขนส่ง"),
            genre=BaseDistributorGenre.objects.create(id="19", name="ค่าจ้างบรรทุก-บุคคลภายนอก"),
        )
        Distributor.objects.create(id="T02", name="ร้านไม่มีชนิด", address=None, tel="")
        self.login(self.user)
        response = self.client.get(reverse('viewVendorReport'))
        self.assertContains(response, "ชนิดผู้จัดจำหน่าย")
        self.assertContains(response, "ประเภทผู้จัดจำหน่าย")
        self.assertContains(response, "<td>ขนส่ง</td>")
        self.assertContains(response, "<td>ค่าจ้างบรรทุก-บุคคลภายนอก</td>")
        self.assertContains(response, "<td>-</td>")
        # T02 ไม่มีที่อยู่ / เบอร์ / ชนิด / ประเภท -> "-" ทั้ง 4 ช่อง ไม่มี "None"
        row = response.content.decode().split('<th scope="row">T02</th>')[1].split("</tr>")[0]
        self.assertEqual(row.count("<td>-</td>"), 4)
        self.assertNotIn("None", row)

    @patch("stock.views.fetch_vendors_from_openstreetmap", return_value=[])
    def test_report_region_presets(self, fetch):
        self.login(self.user)
        response = self.client.get(reverse('viewExternalVendorReport'))
        self.assertContains(response, 'data-lat="9.138200" data-lng="99.321700"')
        self.assertIsNone(response.context["region"])

        # ภาคที่เลือกค้างไว้ในหน้า + ค้นตามพิกัดของภาคนั้น
        response = self.client.get(reverse('viewExternalVendorReport') + "?region=south&lat=9.138200&lng=99.321700&source=osm")
        self.assertEqual(response.context["region"], "south")
        self.assertContains(response, '<option value="south" data-lat="9.138200" data-lng="99.321700" selected>', html=False)
        self.assertEqual(fetch.call_args[0][:2], (9.1382, 99.3217))

        # ค่าที่ไม่รู้จักถูกตัดทิ้ง
        response = self.client.get(reverse('viewExternalVendorReport') + "?region=mars")
        self.assertIsNone(response.context["region"])

    def test_form_page_renders_rules(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        response = self.client.get(reverse('distributorCandidateForm', args=[c.pk]))
        self.assertContains(response, "ความถูกต้องทางกฎหมาย")
        self.assertContains(response, f'name="ans_{self.rules["9.1"].pk}"')
        self.assertNotContains(response, f'name="ans_{self.rules["9"].pk}"')

    def test_requester_detail_page_is_read_only_even_for_approver(self):
        c = self.submitted_candidate()
        for user in (self.user, self.approver):
            self.login(user)
            response = self.client.get(reverse('distributorCandidateDetail', args=[c.pk]))
            self.assertContains(response, OSM_SHOP["name"])
            self.assertContains(response, "สรุป: ผ่านเกณฑ์บังคับ")
            self.assertNotContains(response, "อนุมัติ + สร้างใหม่")

    def test_approve_pages_only_for_group(self):
        c = self.submitted_candidate(answers={"1": "fail"})
        self.login(self.approver)
        response = self.client.get(reverse('distributorApproveList'))
        self.assertEqual([x.pk for x in response.context["candidates"]], [c.pk])
        self.assertContains(response, "ไม่ผ่าน")
        self.assertContains(response, "ข้อ 1 ความถูกต้องทางกฎหมาย")
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertContains(response, "อนุมัติ + สร้างใหม่")
        self.assertContains(response, "ชนิดภาษี")

        self.login(self.user)
        for url in (reverse('distributorApproveList'), reverse('distributorApproveDetail', args=[c.pk])):
            self.assertRedirects(self.client.get(url), reverse('distributorCandidateList'), fetch_redirect_response=False)

    def test_approve_returns_to_queue(self):
        c = self.submitted_candidate()
        response = self.approve(c, mode="new", distributor_id="V900")
        self.assertRedirects(response, reverse('distributorApproveList'), fetch_redirect_response=False)

    def test_sidebar_approve_menu_and_badge_only_for_approver(self):
        self.submitted_candidate()
        approve_url = reverse('distributorApproveList')
        # sidebar อยู่ทุกหน้า: เช็คทั้งหน้าในกลุ่มผู้จัดจำหน่าย และหน้าอื่น
        for name in ('viewVendorReport', 'home'):
            self.login(self.approver)
            response = self.client.get(reverse(name))
            self.assertEqual(response.context["dist_ap_count"], 1)
            self.assertContains(response, f'href="{approve_url}"')

            self.login(self.user)
            response = self.client.get(reverse(name))
            self.assertNotIn("dist_ap_count", response.context)
            self.assertNotContains(response, f'href="{approve_url}"')

    @patch("stock.views.fetch_vendors_from_tomtom")
    @patch("stock.views.fetch_vendors_from_openstreetmap")
    def test_vendor_pages_render_and_highlight_menu(self, fetch_osm, fetch_tomtom):
        self.login(self.approver)
        pages = {
            'viewVendorReport': "vd_in_page",
            'viewExternalVendorReport': "vd_ext_page",
            'distributorCandidateList': "vd_cand_page",
            'distributorApproveList': "vd_ap_page",
        }
        for name, flag in pages.items():
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 200, name)
            self.assertEqual(response.context[flag], "tab-active", name)
            self.assertEqual(response.context["vd_show"], "show", name)
            for other in pages:
                self.assertContains(response, f'href="{reverse(other)}"')
        # หน้าในระบบ / ใบขอเพิ่ม ไม่ต้องค้นแผนที่ / หน้าค้นหาไม่มีพิกัด ก็ยังไม่ยิง API
        fetch_osm.assert_not_called()
        fetch_tomtom.assert_not_called()

    def test_internal_page_only_lists_distributors(self):
        Distributor.objects.create(id="V001", name="ร้านหนึ่ง")
        Distributor.objects.create(id="V002", name="ร้านสอง")
        self.login(self.user)
        response = self.client.get(reverse('viewVendorReport') + "?distributor_id_from=V002&distributor_id_to=V002")
        self.assertEqual([v.id for v in response.context["vender"]], ["V002"])
        self.assertNotIn("external_vendors", response.context)
        self.assertNotIn("candidates", response.context)

    def test_cancel_draft_returns_to_candidate_list(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        response = self.client.post(reverse('distributorCandidateForm', args=[c.pk]), {"action": "cancel"})
        self.assertRedirects(response, reverse('distributorCandidateList'), fetch_redirect_response=False)
        self.assertFalse(DistributorCandidate.objects.exists())

    def test_print_page_shows_signers_and_supplier_code(self):
        c = self.submitted_candidate()
        self.approve(c, mode="new", distributor_id="V900")
        self.login(self.user)
        response = self.client.get(reverse('distributorCandidatePrint', args=[c.pk]))
        self.assertContains(response, "FM-PU-005 Rev.01")
        self.assertContains(response, "V900")
        self.assertContains(response, "requester")
        self.assertContains(response, "approver")

    def test_pages_render_when_users_deleted(self):
        # ผู้ใช้ถูกลบ -> requested_by / approved_by เป็น NULL หน้าเว็บต้องไม่พัง
        c = self.submitted_candidate()
        self.approve(c, mode="new", distributor_id="V900")
        DistributorCandidate.objects.filter(pk=c.pk).update(requested_by=None, approved_by=None)
        self.login(self.approver)
        for name in ('distributorCandidateDetail', 'distributorApproveDetail', 'distributorCandidatePrint'):
            self.assertEqual(self.client.get(reverse(name, args=[c.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('distributorApproveList') + "?status=approved").status_code, 200)
        self.assertEqual(self.client.get(reverse('distributorCandidateList') + "?cand_status=approved").status_code, 200)

    # ----- เพิ่มร้านเอง (ไม่เจอในแผนที่) -----

    def manual_post(self, action="save", answers=None, skip=(), **header):
        data = {"action": action, "name": "หจก. ร้านกรอกเอง", "address": "1 ถ.ทดสอบ", "tel": "081-111-1111"}
        data.update(header)
        answers = answers or {}
        for code in LEAF_CODES:
            if code not in skip:
                data[f"ans_{self.rules[code].pk}"] = answers.get(code, "pass")
        self.login(self.user)
        return self.client.post(reverse('distributorCandidateNew'), data)

    def test_manual_form_page_renders_same_form_without_saving(self):
        self.login(self.user)
        response = self.client.get(reverse('distributorCandidateNew'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ร้านที่กรอกเอง")
        self.assertContains(response, f'name="ans_{self.rules["9.1"].pk}"')
        self.assertContains(response, reverse('distributorCandidateNewCheck'))
        self.assertNotContains(response, 'name="lat"')
        self.assertFalse(DistributorCandidate.objects.exists())

    def test_manual_save_creates_draft_and_goes_to_normal_form(self):
        response = self.manual_post(skip=("11",))
        c = DistributorCandidate.objects.get()
        self.assertRedirects(response, reverse('distributorCandidateForm', args=[c.pk]), fetch_redirect_response=False)
        self.assertEqual((c.place_source, c.status, c.requested_by), ("manual", DistributorCandidate.STATUS_DRAFT, self.user))
        self.assertTrue(c.place_id.startswith("manual-"))
        self.assertEqual((c.name, c.tel), ("หจก. ร้านกรอกเอง", "081-111-1111"))
        self.assertIsNone(c.latitude)
        self.assertEqual(c.answers.count(), len(LEAF_CODES) - 1)
        # หน้าฟอร์มเดิมเปิดต่อได้ / ใบที่ 2 ได้ place_id ไม่ซ้ำ
        self.assertEqual(self.client.get(reverse('distributorCandidateForm', args=[c.pk])).status_code, 200)
        self.manual_post(name="ร้านที่สอง")
        self.assertEqual(DistributorCandidate.objects.values("place_id").distinct().count(), 2)

    def test_manual_submit_then_approve_into_distributor(self):
        self.manual_post(action="submit")
        c = DistributorCandidate.objects.get()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.approve(c, mode="new", distributor_id="V950")
        d = Distributor.objects.get(id="V950")
        self.assertEqual((d.name, d.place_source), ("หจก. ร้านกรอกเอง", "manual"))

    def test_manual_requires_name_and_keeps_typed_values(self):
        response = self.manual_post(name="  ", tel="099-999-9999")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(DistributorCandidate.objects.exists())
        self.assertContains(response, "กรุณากรอกชื่อบริษัท/ร้านค้า")
        self.assertContains(response, 'value="099-999-9999"')

    def test_manual_cancel_saves_nothing(self):
        response = self.manual_post(action="cancel")
        self.assertRedirects(response, reverse('distributorCandidateList'), fetch_redirect_response=False)
        self.assertFalse(DistributorCandidate.objects.exists())

    def test_manual_check_works_before_saving(self):
        self.login(self.user)
        data = {f"ans_{self.rules[code].pk}": "pass" for code in LEAF_CODES}
        response = self.client.post(reverse('distributorCandidateNewCheck'), data)
        self.assertEqual(response.json()["passed"], True)
        self.assertFalse(DistributorCandidate.objects.exists())

    def test_manual_warns_on_similar_name(self):
        Distributor.objects.create(id="V001", name="บริษัท ร้านกรอกเอง จำกัด")
        self.manual_post()
        self.assertTrue(DistributorCandidate.objects.exists())
        warnings = [str(m) for m in self.client.get(reverse('distributorCandidateList')).context["messages"]]
        self.assertTrue(any("V001" in w for w in warnings), warnings)

    def test_add_manual_buttons_on_list_and_search_pages(self):
        self.login(self.user)
        for name in ('distributorCandidateList', 'viewExternalVendorReport'):
            self.assertContains(self.client.get(reverse(name)), reverse('distributorCandidateNew'))

    # ----- คำนำหน้า / เครดิต (ตาราง BasePrefix / BaseCredit ของ Express) -----

    def test_prefix_and_credit_dropdowns_saved_shown_and_copied_on_approve(self):
        from stock.models import BaseCredit, BasePrefix
        hjk = BasePrefix.objects.create(id="01", name="หจก.")
        BasePrefix.objects.create(id="03", name="บจก.")
        c30 = BaseCredit.objects.create(id=3, name="30 วัน")
        BaseCredit.objects.create(id=14, name="เข้าลิสซิ่ง")
        BaseCredit.objects.create(id=15, name="7 วัน")

        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.user)
        response = self.client.get(reverse('distributorCandidateForm', args=[c.pk]))
        self.assertContains(response, 'name="prefix"')
        # เรียงตามจำนวนวัน ค่าที่ไม่ใช่ตัวเลขอยู่ท้าย
        self.assertEqual([x.name for x in response.context["credit_options"]], ["7 วัน", "30 วัน", "เข้าลิสซิ่ง"])

        self.fill_form(c, action="submit", prefix=hjk.pk, credit=c30.pk)
        c.refresh_from_db()
        self.assertEqual((c.prefix, c.credit), (hjk, c30))
        for name in ('distributorCandidateDetail', 'distributorCandidatePrint'):
            response = self.client.get(reverse(name, args=[c.pk]))
            self.assertContains(response, f"หจก. {OSM_SHOP['name']}")
            self.assertContains(response, "30 วัน")

        self.approve(c, mode="new", distributor_id="V901")
        d = Distributor.objects.get(id="V901")
        self.assertEqual((d.prefix, d.credit), (hjk, c30))

    def test_unknown_prefix_or_credit_is_ignored(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.fill_form(c, prefix="zz", credit="999")
        c.refresh_from_db()
        self.assertEqual((c.prefix, c.credit), (None, None))

    def test_other_contact_shown_and_exported(self):
        import io as _io
        import openpyxl
        c = self.submitted_candidate(other_contact="LINE: @shop\nFacebook: ร้านทดสอบ")
        for name in ('distributorCandidateDetail', 'distributorCandidatePrint'):
            self.assertContains(self.client.get(reverse(name, args=[c.pk])), "LINE: @shop<br>Facebook: ร้านทดสอบ")

        self.approve(c, mode="new", distributor_id="V902")
        self.login(self.user)
        ws = openpyxl.load_workbook(_io.BytesIO(self.client.get(reverse('exportExcelVendor')).content)).active
        header = [cell.value for cell in ws[1]]
        row = {h: cell.value for h, cell in zip(header, ws[2])}
        self.assertEqual(header.index("ช่องทางติดต่ออื่นๆ"), header.index("แฟกซ์") + 1)
        self.assertEqual(row["ช่องทางติดต่ออื่นๆ"], "LINE: @shop\nFacebook: ร้านทดสอบ")

    def test_business_type_not_in_form_and_kept_when_saving(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        DistributorCandidate.objects.filter(pk=c.pk).update(business_type="ค่าเดิม")
        self.login(self.user)
        self.assertNotContains(self.client.get(reverse('distributorCandidateForm', args=[c.pk])), 'name="business_type"')
        self.fill_form(c, contact="คุณสมชาย")
        c.refresh_from_db()
        self.assertEqual((c.contact, c.business_type), ("คุณสมชาย", "ค่าเดิม"))

    def test_vendor_excel_export_keeps_formula_like_values_as_text(self):
        import io as _io
        import openpyxl
        Distributor.objects.create(id="V778", name="=HYPERLINK(\"http://x\",\"x\")", other_contact="=1+1")
        self.login(self.user)
        ws = openpyxl.load_workbook(_io.BytesIO(self.client.get(reverse('exportExcelVendor')).content)).active
        header = [cell.value for cell in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "V778")
        for column in ("ชื่อผู้จัดจำหน่าย", "ช่องทางติดต่ออื่นๆ"):
            cell = row[header.index(column)]
            self.assertEqual(cell.data_type, "s", column)
        self.assertEqual(row[header.index("ช่องทางติดต่ออื่นๆ")].value, "=1+1")

    # ----- สิทธิ์เมนู "ผู้จัดจำหน่าย" ตามกลุ่มผู้ใช้ -----

    VENDOR_PAGES = ('viewVendorReport', 'viewExternalVendorReport', 'distributorCandidateList', 'distributorCandidateNew', 'exportExcelVendor')

    def test_distributor_tab_created_by_migration(self):
        self.assertTrue(BaseVisible.objects.filter(name='Distributor').exists())
        self.assertFalse(Group.objects.filter(name='Distributor').exists())

    def test_outsider_has_no_vendor_menu_and_pages_redirect_home(self):
        self.propose()
        c = DistributorCandidate.objects.get()
        self.login(self.outsider)
        response = self.client.get(reverse('home'))
        self.assertFalse(response.context["can_vendor_menu"])
        self.assertNotContains(response, 'id="vendorSubmenu"')
        for name in self.VENDOR_PAGES:
            self.assertRedirects(self.client.get(reverse(name)), reverse('home'), fetch_redirect_response=False, msg_prefix=name)
        for name in ('distributorCandidateDetail', 'distributorCandidateForm', 'distributorCandidatePrint', 'distributorApproveDetail'):
            self.assertRedirects(self.client.get(reverse(name, args=[c.pk])), reverse('home'), fetch_redirect_response=False, msg_prefix=name)
        self.assertRedirects(self.client.get(reverse('distributorApproveList')), reverse('home'), fetch_redirect_response=False)
        # เสนอเพิ่ม / กรอกเอง ผ่าน POST ก็ไม่ได้
        before = DistributorCandidate.objects.count()
        self.client.post(reverse('distributorCandidateNew'), {"action": "save", "name": "ร้านแอบเพิ่ม"})
        self.client.post(reverse('proposeDistributorCandidate'), {"place_source": "osm", "place_id": "node/9", "name": "x"})
        self.assertEqual(DistributorCandidate.objects.count(), before)

    def test_distributor_group_sees_vendor_menu_without_approve(self):
        self.login(self.user)
        response = self.client.get(reverse('home'))
        self.assertTrue(response.context["can_vendor_menu"])
        for name in ('viewVendorReport', 'viewExternalVendorReport', 'distributorCandidateList'):
            self.assertContains(response, f'href="{reverse(name)}"')
        self.assertNotContains(response, f'href="{reverse("distributorApproveList")}"')
        for name in self.VENDOR_PAGES:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        self.assertRedirects(self.client.get(reverse('distributorApproveList')), reverse('distributorCandidateList'), fetch_redirect_response=False)

    def test_approve_group_alone_sees_everything(self):
        self.assertFalse(self.approver.userprofile.visible.filter(name='Distributor').exists())
        self.login(self.approver)
        response = self.client.get(reverse('home'))
        for name in ('viewVendorReport', 'viewExternalVendorReport', 'distributorCandidateList', 'distributorApproveList'):
            self.assertContains(response, f'href="{reverse(name)}"')
        for name in self.VENDOR_PAGES + ('distributorApproveList',):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    # ----- รหัส Supplier อัตโนมัติ (XX Y ZZZ) -----

    def test_supplier_code_letter_rules(self):
        from stock.views import supplier_code_letter
        cases = {
            "กิจเจริญ": "ก",
            "เกื้อเจริญเครื่องครัว": "ก",    # ข้ามสระหน้า
            "โก้จุ้ย": "ก",
            "ร้านไพศาลเมืองไฟฟ้า": "พ",      # ข้ามคำนำหน้าร้าน
            "บริษัท ศรีสุทธา จำกัด": "ศ",
            "ห้างหุ้นส่วนจำกัด แหลมทองฮาร์ดแวร์": "ห",
            "หจก. ติยะฮาร์ดแวร์": "ต",
            "นาย\xa0วีรยุทร": "น",            # คำนำหน้าบุคคลไม่ตัด (ตามรหัสเดิม)
            "ABC Shop": "อ",                  # ชื่ออังกฤษ: แปลงตามเสียงอ่าน
            "Siam Global": "ส",
            "bangkok steel": "บ",
            "Computer Plus": "ค",
            "3K Shop": None,
            "": None,
        }
        for name, letter in cases.items():
            self.assertEqual(supplier_code_letter(name), letter, name)

    def test_suggest_distributor_id_runs_per_type_and_letter(self):
        from stock.models import BaseDistributorType
        from stock.views import suggest_distributor_id
        BaseDistributorType.objects.create(id="01", name="ทั่วไป")
        for code in ("01ก001", "01ก002", "01ข007", "03ก010", "01ก0001", "01ก"):
            Distributor.objects.create(id=code, name="x")
        self.assertEqual(suggest_distributor_id("01", "กิจเจริญ"), "01ก003")
        self.assertEqual(suggest_distributor_id("01", "ขวัญใจ"), "01ข008")
        self.assertEqual(suggest_distributor_id("03", "เก่งการค้า"), "03ก011")
        self.assertEqual(suggest_distributor_id("02", "จันทร์"), "02จ001")
        self.assertEqual(suggest_distributor_id(None, "กิจเจริญ"), "01ก003")  # ไม่มีชนิด -> 01
        self.assertEqual(suggest_distributor_id("01", "ABC"), "01อ001")
        self.assertIsNone(suggest_distributor_id("01", "3K Shop"))
        # เลข 999 ถูกใช้ไปแล้ว -> ใช้เลขว่างตัวแรก
        Distributor.objects.create(id="01ข999", name="x")
        self.assertEqual(suggest_distributor_id("01", "ขวัญใจ"), "01ข001")

    def test_approve_page_prefills_suggested_supplier_code(self):
        from stock.models import BaseDistributorType
        general = BaseDistributorType.objects.create(id="01", name="ทั่วไป")
        Distributor.objects.create(id="01จ004", name="จันทร์เจริญ")  # ร้านในใบ = "ร้านเจริญการช่าง" -> จ
        c = self.submitted_candidate(type=general.pk)
        self.assertEqual(c.type, general)
        self.login(self.approver)
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertEqual(response.context["suggested_distributor_id"], "01จ005")
        self.assertContains(response, 'value="01จ005"')

        # ไม่มีชนิด -> XX = 01
        DistributorCandidate.objects.filter(pk=c.pk).update(type=None)
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertEqual(response.context["suggested_distributor_id"], "01จ005")
        # คำอธิบายรูปแบบรหัส / คำเตือนชื่อภาษาอังกฤษ ไม่แสดงบนหน้าอนุมัติ
        self.assertNotContains(response, "ใบนี้ไม่ได้เลือกชนิดผู้จัดจำหน่าย ใช้ 01")
        self.assertNotContains(response, "รูปแบบ XX Y ZZZ")

        # ชื่อภาษาอังกฤษ -> เดาตัวอักษรไทยตามเสียงอ่าน
        DistributorCandidate.objects.filter(pk=c.pk).update(name="Siam Global")
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertEqual(response.context["suggested_distributor_id"], "01ส001")
        self.assertNotContains(response, "ระบบเดาตัวอักษรไทยตามเสียงอ่าน")

        # ชื่อขึ้นต้นด้วยตัวเลข -> ว่างไว้ให้กรอกเอง
        DistributorCandidate.objects.filter(pk=c.pk).update(name="3K Shop")
        response = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertIsNone(response.context["suggested_distributor_id"])
        self.assertContains(response, "ตั้งให้อัตโนมัติไม่ได้")

    # ----- ชื่อบริษัท + สาขา -----

    def test_name_with_branch_rules(self):
        from stock.views import name_with_branch
        self.assertEqual(name_with_branch("หจก. ร้านกรอกเอง", ""), "หจก. ร้านกรอกเอง")
        self.assertEqual(name_with_branch("  ", "2"), "")  # ชื่อว่าง ต้องยังว่าง ไม่งั้นผ่านเช็คช่องบังคับ
        self.assertEqual(name_with_branch("หจก. ร้านกรอกเอง", " 2 "), "หจก. ร้านกรอกเอง (สาขา 2)")
        # บันทึกซ้ำไม่ต่อซ้อน / เปลี่ยนสาขา / ลบสาขา
        self.assertEqual(name_with_branch("หจก. ร้านกรอกเอง (สาขา 2)", "2"), "หจก. ร้านกรอกเอง (สาขา 2)")
        self.assertEqual(name_with_branch("หจก. ร้านกรอกเอง (สาขา 2)", "00003"), "หจก. ร้านกรอกเอง (สาขา 00003)")
        self.assertEqual(name_with_branch("หจก. ร้านกรอกเอง (สาขา 2)", ""), "หจก. ร้านกรอกเอง")
        # กรอกคำว่า "สาขา" มาเอง / สำนักงานใหญ่
        self.assertEqual(name_with_branch("ร้าน ก", "สาขาบางนา"), "ร้าน ก (สาขาบางนา)")
        self.assertEqual(name_with_branch("ร้าน ก", "สำนักงานใหญ่"), "ร้าน ก (สำนักงานใหญ่)")
        self.assertEqual(name_with_branch("ร้าน ก (สำนักงานใหญ่)", "2"), "ร้าน ก (สาขา 2)")
        # วงเล็บอื่นในชื่อร้านไม่แตะ
        self.assertEqual(name_with_branch("บริษัท เอ (ประเทศไทย) จำกัด", "2"), "บริษัท เอ (ประเทศไทย) จำกัด (สาขา 2)")
        # ยาวเกินช่อง -> ตัดชื่อ แต่สาขายังอยู่ครบ
        long_name = name_with_branch("ก" * 255, "2")
        self.assertEqual(len(long_name), 255)
        self.assertTrue(long_name.endswith(" (สาขา 2)"))
        # POST ยาวผิดปกติ (ช่องว่างกลางชื่อเป็นแสนตัว) ต้องไม่ค้าง และได้ไม่เกินความยาวช่อง
        import time
        started = time.monotonic()
        huge = name_with_branch("ก" + " " * 200000 + "ข", "2" * 300000)
        self.assertLess(time.monotonic() - started, 1)
        self.assertLessEqual(len(huge), 255)

    def test_manual_form_appends_branch_to_name(self):
        response = self.manual_post(name="  ", branch="2")
        self.assertContains(response, "กรุณากรอกชื่อบริษัท/ร้านค้า")
        self.assertFalse(DistributorCandidate.objects.exists())

        self.manual_post(branch="2")
        c = DistributorCandidate.objects.get()
        self.assertEqual((c.name, c.branch), ("หจก. ร้านกรอกเอง (สาขา 2)", "2"))
        # เปิดฟอร์มเดิมแล้วบันทึกซ้ำ ชื่อไม่ต่อซ้อน / เปลี่ยนสาขาแล้วชื่อเปลี่ยนตาม
        self.fill_form(c)
        c.refresh_from_db()
        self.assertEqual(c.name, "หจก. ร้านกรอกเอง (สาขา 2)")
        self.fill_form(c, branch="5")
        c.refresh_from_db()
        self.assertEqual(c.name, "หจก. ร้านกรอกเอง (สาขา 5)")
        self.fill_form(c, branch="")
        c.refresh_from_db()
        self.assertEqual(c.name, "หจก. ร้านกรอกเอง")

    def test_form_page_has_name_preview(self):
        self.login(self.user)
        response = self.client.get(reverse('distributorCandidateNew'))
        self.assertContains(response, 'id="f-name-preview"')

    # ----- check รหัส Supplier ซ้ำอีกรอบตอนกดอนุมัติ -----

    def test_approve_renumbers_suggested_code_taken_meanwhile(self):
        from django.contrib.messages import get_messages
        c = self.submitted_candidate()  # "ร้านเจริญการช่าง" ไม่มีชนิด -> 01จ
        self.login(self.approver)
        page = self.client.get(reverse('distributorApproveDetail', args=[c.pk]))
        self.assertEqual(page.context["suggested_distributor_id"], "01จ001")
        self.assertContains(page, 'name="suggested_distributor_id" value="01จ001"')
        # เปิดหน้าค้างไว้ ระหว่างนั้นมีร้านอื่นได้รหัส 01จ001 ไปก่อน
        Distributor.objects.create(id="01จ001", name="จันทร์เจริญ")
        response = self.approve(c, mode="new", distributor_id="01จ001", suggested_distributor_id="01จ001")
        c.refresh_from_db()
        self.assertEqual((c.status, c.distributor_id), (DistributorCandidate.STATUS_APPROVED, "01จ002"))
        self.assertEqual(Distributor.objects.get(id="01จ002").name, c.name)
        self.assertEqual(Distributor.objects.get(id="01จ001").name, "จันทร์เจริญ")  # ร้านเดิมไม่โดนเขียนทับ
        text = " ".join(str(m) for m in get_messages(response.wsgi_request))
        self.assertIn("รหัส 01จ001 ถูกใช้ไปแล้วก่อนกดอนุมัติ ระบบรันเลขใหม่ให้", text)

    def test_approve_typed_code_taken_is_not_renumbered(self):
        c = self.submitted_candidate()
        Distributor.objects.create(id="01จ005", name="ร้านอื่น")
        # ผู้อนุมัติแก้รหัสเอง (ไม่ตรงกับที่ระบบตั้งให้) -> ไม่เดาให้ เด้งกลับให้แก้
        self.approve(c, mode="new", distributor_id="01จ005", suggested_distributor_id="01จ001")
        c.refresh_from_db()
        self.assertEqual(c.status, DistributorCandidate.STATUS_SUBMITTED)
        self.assertFalse(Distributor.objects.filter(id="01จ006").exists())
        self.assertEqual(Distributor.objects.get(id="01จ005").name, "ร้านอื่น")

    # ----- ต้นทางค้นหาผู้จัดจำหน่าย (เก็บใน DB เพิ่ม/ลบในหน้า admin) -----

    def test_vendor_origins_seeded_by_migration(self):
        from stock.models import BaseVendorOrigin
        origins = list(BaseVendorOrigin.objects.all())
        self.assertEqual(len(origins), 17)
        self.assertEqual(origins[0].name, "ศิลาชัย 1")
        self.assertEqual(origins[-1].name, "เหมืองคลองปราบ")
        first = origins[0]
        self.assertEqual((first.latitude, first.longitude), (Decimal("8.935915"), Decimal("98.960906")))
        self.assertEqual(first.map_url, "https://maps.app.goo.gl/YhmoZ5KzpFvuNKvc9")

    @patch("stock.views.fetch_vendors_from_openstreetmap", return_value=[])
    def test_report_origin_presets(self, fetch):
        from stock.models import BaseVendorOrigin
        site = BaseVendorOrigin.objects.get(name="กงตาก 3")
        BaseVendorOrigin.objects.filter(name="เหมืองคลองปราบ").update(is_active=False)
        self.login(self.user)
        response = self.client.get(reverse('viewExternalVendorReport'))
        self.assertContains(response, f'<option value="{site.pk}" data-lat="8.956197" data-lng="99.471246" >กงตาก 3</option>', html=False)
        self.assertNotContains(response, "เหมืองคลองปราบ")  # ปิดใช้งานแล้วไม่แสดง
        self.assertIsNone(response.context["origin"])

        # ต้นทางที่เลือกค้างไว้ในหน้า + ค้นตามพิกัดของต้นทางนั้น
        response = self.client.get(reverse('viewExternalVendorReport') + f"?origin={site.pk}&lat=8.956197&lng=99.471246&source=osm")
        self.assertEqual(response.context["origin"], site)
        self.assertContains(response, f'<option value="{site.pk}" data-lat="8.956197" data-lng="99.471246" selected>', html=False)
        self.assertContains(response, "ต้นทาง: กงตาก 3")
        self.assertEqual(fetch.call_args[0][:2], (8.956197, 99.471246))

        # ค่าที่ไม่รู้จักถูกตัดทิ้ง
        response = self.client.get(reverse('viewExternalVendorReport') + "?origin=99999")
        self.assertIsNone(response.context["origin"])

    def test_vendor_origin_admin_accepts_pasted_coordinates(self):
        from stock.admin import BaseVendorOriginAdminForm
        form = BaseVendorOriginAdminForm(data={
            "name": "ไซต์ใหม่", "coordinates": " 8.935915423841209 , 98.96090612019505 ", "order": 1, "is_active": True,
        })
        self.assertTrue(form.is_valid(), form.errors)
        origin = form.save()
        self.assertEqual((origin.latitude, origin.longitude), (Decimal("8.935915"), Decimal("98.960906")))
        # เปิดแก้ไขแล้วเห็นพิกัดเดิมในช่องเดียว
        self.assertEqual(BaseVendorOriginAdminForm(instance=origin).fields["coordinates"].initial, "8.935915, 98.960906")

        for bad in ("8.93", "abc, 98.9", "8.9, 98.9, 1", "91, 98.9", "8.9, 181", "NaN, 98.9"):
            form = BaseVendorOriginAdminForm(data={"name": "x", "coordinates": bad, "order": 1})
            self.assertFalse(form.is_valid(), bad)
            self.assertIn("coordinates", form.errors)
