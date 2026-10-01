from io import BytesIO

import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User

from stock.models import (
    UserProfile, BaseBranchCompany, BaseAddress, BranchCompanyBaseAdress, CarLogbook,
    BaseCar,
)


class CarLogbookAllCompanyTestCase(TestCase):
    """/carLogBook/ และ /report/carLogBook/ ในแท็ป ALL ต้องแสดงใบบันทึกการใช้รถ
    ของทุกบริษัทที่ user มีสิทธิ์มองเห็น ส่วนแท็ปบริษัทยังแสดงเฉพาะบริษัทนั้น"""

    URL_NAMES = ("viewCL", "viewCLReport")

    @classmethod
    def setUpTestData(cls):
        cls.all_tab = BaseBranchCompany.objects.create(id="0", code="ALL", name="All")
        cls.ho = BaseBranchCompany.objects.create(id="1", code="HO", name="Head Office")
        cls.br = BaseBranchCompany.objects.create(id="2", code="BR", name="Branch")
        cls.hidden = BaseBranchCompany.objects.create(id="3", code="XX", name="Hidden")
        address = BaseAddress.objects.create(name_th="Test Co", address="1 Test Rd")
        for b in (cls.ho, cls.br, cls.hidden):
            BranchCompanyBaseAdress.objects.create(branch_company=b, address=address)

        cls.user = User.objects.create_user(username="cl_user", password="pw")
        profile = UserProfile.objects.create(user=cls.user)
        profile.branch_company.add(cls.all_tab, cls.ho, cls.br)

        cls.cl_ho = CarLogbook.objects.create(branch_company=cls.ho, ref_no="CL-HO-1")
        cls.cl_br = CarLogbook.objects.create(branch_company=cls.br, ref_no="CL-BR-1")
        cls.cl_hidden = CarLogbook.objects.create(branch_company=cls.hidden, ref_no="CL-XX-1")

    def _get(self, url_name, company_code):
        client = Client()
        client.login(username="cl_user", password="pw")
        session = client.session
        session["company_code"] = company_code
        session.save()
        return client.get(reverse(url_name))

    def test_all_tab_shows_every_permitted_company(self):
        for url_name in self.URL_NAMES:
            with self.subTest(url_name=url_name):
                response = self._get(url_name, "ALL")
                self.assertEqual(response.status_code, 200)
                ids = {cl.id for cl in response.context["cls"]}
                self.assertEqual(ids, {self.cl_ho.id, self.cl_br.id})
                self.assertTrue(response.context["is_all_comp"])
                self.assertContains(response, "<th scope=\"col\">บริษัท</th>", html=False)

    def test_company_tab_shows_only_that_company(self):
        for url_name in self.URL_NAMES:
            with self.subTest(url_name=url_name):
                response = self._get(url_name, "HO")
                self.assertEqual(response.status_code, 200)
                ids = {cl.id for cl in response.context["cls"]}
                self.assertEqual(ids, {self.cl_ho.id})
                self.assertFalse(response.context["is_all_comp"])
                self.assertNotContains(response, "<th scope=\"col\">บริษัท</th>", html=False)


class CarLogbookExcelCompanyNameTestCase(TestCase):
    """Excel รายงานบันทึกการใช้รถประจำวัน และ สรุปค่าใช้จ่ายแต่ละหน่วยงาน
    ต้องแสดงชื่อบริษัทของใบบันทึกการใช้รถ (แท็ป ALL มีหลายบริษัท)"""

    @classmethod
    def setUpTestData(cls):
        # ไม่ตั้ง invoice_code เพื่อไม่ให้ไปดึงค่าอะไหล่จาก pg_db
        cls.all_tab = BaseBranchCompany.objects.create(id="0", code="ALL", name="All")
        cls.ho = BaseBranchCompany.objects.create(id="1", code="HO", name="Head Office")
        cls.br = BaseBranchCompany.objects.create(id="2", code="BR", name="Branch")
        address = BaseAddress.objects.create(name_th="Test Co", address="1 Test Rd")
        for b in (cls.ho, cls.br):
            BranchCompanyBaseAdress.objects.create(branch_company=b, address=address)

        cls.user = User.objects.create_user(username="cl_excel", password="pw")
        profile = UserProfile.objects.create(user=cls.user)
        profile.branch_company.add(cls.all_tab, cls.ho, cls.br)

        cls.car_shared = BaseCar.objects.create(code="C1", name="Truck1")
        cls.car_ho = BaseCar.objects.create(code="C2", name="Truck2")
        miles = dict(mile_start_job1=100, mile_end_job1=150)
        CarLogbook.objects.create(branch_company=cls.ho, car=cls.car_shared, ref_no="CL-1", **miles)
        CarLogbook.objects.create(branch_company=cls.br, car=cls.car_shared, ref_no="CL-2", **miles)
        CarLogbook.objects.create(branch_company=cls.ho, car=cls.car_ho, ref_no="CL-3", **miles)

    def _workbook(self, url_name, company_code):
        client = Client()
        client.login(username="cl_excel", password="pw")
        session = client.session
        session["company_code"] = company_code
        session.save()
        response = client.get(reverse(url_name))
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content) if response.streaming else response.content
        return openpyxl.load_workbook(BytesIO(content))

    def test_daily_report_shows_record_company_names(self):
        wb = self._workbook("excelDailyCL", "ALL")
        self.assertEqual(wb["C1 Truck1"]["B2"].value, "Branch, Head Office")
        self.assertEqual(wb["C2 Truck2"]["B2"].value, "Head Office")

    def test_daily_report_company_tab(self):
        wb = self._workbook("excelDailyCL", "HO")
        self.assertEqual(wb["C1 Truck1"]["B2"].value, "Head Office")

    def test_expenses_summary_has_company_column(self):
        ws = self._workbook("excelExpensesByCarLog", "ALL").active
        self.assertEqual(ws["B4"].value, "บริษัท")  # B4:B5 merge
        self.assertEqual(ws["E4"].value, "ปริมาณที่ใช้")  # คอลัมน์อื่นเลื่อนไป 1
        rows = [(r[1], r[2]) for r in ws.iter_rows(min_row=6, max_col=3, values_only=True)]
        self.assertEqual(rows, [("Branch", "C1"), ("Head Office", "C1"), ("Head Office", "C2")])

    def test_expenses_summary_company_tab(self):
        ws = self._workbook("excelExpensesByCarLog", "HO").active
        rows = [(r[1], r[2]) for r in ws.iter_rows(min_row=6, max_col=3, values_only=True)]
        self.assertEqual(rows, [("Head Office", "C1"), ("Head Office", "C2")])
